from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "public" / "downloads" / "密盾传输系统使用手册.docx"
ROOT_COPY = ROOT / "密盾传输系统使用手册.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, enabled=True):
    paragraph.paragraph_format.keep_with_next = enabled


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])
    paragraph.add_run(" 页")


def add_table(doc, headers, rows, widths, compact=False, trailing_space=True):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, (cell, title, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "123B52")
        set_cell_margins(cell, top=75 if compact else 110, bottom=75 if compact else 110)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(10.5)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column_index, (cell, value, width) in enumerate(zip(cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=75 if compact else 110, bottom=75 if compact else 110)
            if row_index % 2:
                set_cell_shading(cell, "F1F7FA")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(10.5)
    if trailing_space:
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_step(doc, number, title, body):
    heading = doc.add_heading(f"{number}  {title}", level=2)
    set_keep_with_next(heading)
    paragraph = doc.add_paragraph(body)
    paragraph.paragraph_format.keep_together = True


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(35, 43, 48)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(6)

    title_style = styles["Title"]
    title_style.font.name = "Microsoft YaHei"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title_style.font.size = Pt(28)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (("Heading 1", 18, 14, 8), ("Heading 2", 14, 10, 5), ("Heading 3", 12, 8, 4)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "密盾传输  文件安全传输系统"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(110, 120, 126)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(110, 120, 126)

    # Cover page
    doc.add_paragraph().paragraph_format.space_after = Pt(90)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("密盾传输系统使用手册")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("基于 AES 256 GCM 与 RSA OAEP 的文件安全传输系统")
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(47, 91, 111)
    doc.add_paragraph().paragraph_format.space_after = Pt(76)

    cover_table = doc.add_table(rows=4, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover_table.autofit = False
    cover_data = [
        ("文档用途", "用户操作与系统演示"),
        ("适用版本", "SecureVault 1.0"),
        ("运行形式", "浏览器网页应用"),
        ("更新时间", "2026 年 9 月"),
    ]
    for row, (label, value) in zip(cover_table.rows, cover_data):
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(3.75)
        for cell in row.cells:
            set_cell_margins(cell, 135, 160, 135, 160)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)
    doc.add_paragraph()
    p = doc.add_paragraph("本手册说明系统的登录、加密上传、授权共享、授权撤销、下载校验、审计查询和维护方法。按照本文步骤操作，即可完成一次可验证、可控制的安全文件传输。")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.55)
    p.paragraph_format.right_indent = Inches(0.55)
    p.paragraph_format.space_before = Pt(22)
    doc.add_page_break()

    # Contents and overview
    doc.add_heading("目录", level=1)
    contents = [
        "1  系统概述", "2  开始使用", "3  加密上传文件", "4  管理与下载文件",
        "5  授权共享文件", "6  接收共享文件", "7  查看安全审计日志",
        "8  加密机制说明", "9  常见问题", "10  答辩演示建议", "11  管理员部署说明",
    ]
    for item in contents:
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(4)

    doc.add_heading("1  系统概述", level=1)
    doc.add_paragraph("密盾传输是一套前后端一体的网页文件安全传输系统。系统把文件内容加密、密钥封装、身份验证、访问控制、完整性校验和审计记录组合成一条完整流程。用户不需要安装桌面客户端，只需使用现代浏览器访问系统入口。")
    doc.add_paragraph("系统的主要结论是：存储层只保存密文，下载请求必须通过服务端身份和共享权限校验，解密后还需核对 SHA 256 摘要。任何一步失败，系统都会拒绝返回原始文件并记录失败事件。")
    doc.add_page_break()
    add_table(doc, ["模块", "用户可完成的任务", "安全控制"], [
        ("控制台", "查看文件、共享和存储统计", "只显示当前账号数据"),
        ("我的文件", "上传 下载 共享 撤销 删除", "所有者身份校验"),
        ("共享给我", "下载他人授权的文件", "邮箱匹配和有效期校验"),
        ("安全日志", "查看操作账号 结果和时间", "本人及所属文件范围"),
        ("使用手册", "在线阅读或下载本文档", "登录后访问"),
    ], [1.15, 2.85, 2.25])

    doc.add_heading("2  开始使用", level=1)
    add_step(doc, "2.1", "准备运行环境", "推荐使用 Chrome、Edge 或 Firefox，浏览器须启用 JavaScript 和 Cookie。演示版单文件最大 15 MB、单账户最多 5 GB，支持 PDF、Word、Excel、图片、压缩包和纯文本等常见类型。")
    add_step(doc, "2.2", "进入系统", "打开老师或管理员提供的网站地址。若当前没有登录，页面会显示登录入口。点击“使用 ChatGPT 账号登录”，完成身份验证后自动返回控制台。")
    add_step(doc, "2.3", "确认账号", "登录后，页面右上角会显示当前用户名称和邮箱。共享权限依赖邮箱精确匹配，因此向他人共享前，应确认对方用于登录的邮箱地址。")
    add_bullet(doc, "控制台显示“系统安全”表示前端已进入已认证工作区。")
    add_bullet(doc, "退出时使用左侧导航底部的“退出登录”，公共电脑上不要只关闭浏览器标签页。")

    doc.add_heading("3  加密上传文件", level=1)
    add_step(doc, "3.1", "选择文件", "在控制台或“我的文件”页面点击“上传文件”，从本机选择一个非空文件。文件不会以明文形式写入系统存储。")
    add_step(doc, "3.2", "等待加密", "按钮显示“正在加密”期间不要重复点击。服务端为本文件生成独立的 256 位 AES 密钥和 96 位随机 IV，使用 AES 256 GCM 加密文件，再用 RSA OAEP 公钥封装 AES 密钥。")
    add_step(doc, "3.3", "确认结果", "出现“文件已加密并安全保存”后，文件会显示在列表中。文件信息包含大小、上传时间和 SHA 256 摘要前缀，便于识别和核对。")
    add_table(doc, ["结果", "含义", "处理方式"], [
        ("上传成功", "密文和元数据已保存", "可立即下载或共享"),
        ("文件为空", "选择了 0 字节文件", "更换有效文件"),
        ("超过 15 MB", "超过演示版限制", "压缩或拆分后上传"),
        ("密钥配置错误", "服务端 RSA 密钥缺失或格式错误", "联系管理员检查环境变量"),
    ], [1.25, 2.3, 2.7])

    doc.add_heading("4  管理与下载文件", level=1)
    add_step(doc, "4.1", "查看文件", "进入“我的文件”，列表展示当前账号拥有的全部文件。不同账号之间的数据在服务端查询阶段就已隔离。")
    add_step(doc, "4.2", "下载文件", "点击下载图标。服务端先确认所有权或有效共享关系，再从对象存储读取密文，使用 RSA 私钥解封 AES 密钥，并用 AES GCM 解密。随后计算 SHA 256；只有摘要一致时浏览器才收到原始文件。")
    add_step(doc, "4.3", "删除文件", "点击删除图标并确认。系统会删除对象存储中的密文、数据库中的文件元数据以及关联共享权限。删除不可撤销，操作前应确认本地已有必要备份。")
    doc.add_paragraph("完整性失败时的正确行为是拒绝下载，而不是尝试忽略错误。AES GCM 的认证标签和 SHA 256 摘要共同帮助系统发现密文损坏、元数据不匹配或异常篡改。")

    doc.add_heading("5  授权共享文件", level=1)
    add_step(doc, "5.1", "打开共享窗口", "在文件列表中点击共享图标。只有文件所有者能创建或更新共享权限。")
    add_step(doc, "5.2", "填写收件人", "输入对方用于登录系统的完整邮箱地址。系统会规范化邮箱大小写，但不会向不存在的账号发送邮件。")
    add_step(doc, "5.3", "设置有效期", "有效期范围为 1 至 168 小时。默认 24 小时。对同一文件和同一邮箱再次共享时，系统会更新有效期，不会创建重复记录。")
    add_step(doc, "5.4", "完成授权", "点击“确认共享”。成功后，文件会出现在对方的“共享给我”页面。共享行为同时写入安全审计日志。若网站采用私密发布，管理员还需要把收件人加入网站访问名单。")
    add_step(doc, "5.5", "管理或撤销授权", "进入“我的文件”，在“当前有效授权”区域查看文件名、收件邮箱、到期时间和下载次数。点击“撤销”并确认后，该收件人会立即失去下载权限，系统同时记录撤销事件。")
    doc.add_paragraph("共享并不会生成公开下载链接。每次下载都要求收件人处于登录状态，并在服务端实时检查网站访问资格、收件邮箱和到期时间。")

    doc.add_heading("6  接收共享文件", level=1)
    add_step(doc, "6.1", "使用正确账号登录", "收件人必须先获得网站访问资格，再使用共享时填写的邮箱登录。如果使用其他账号，即使知道文件编号也不能下载。")
    add_step(doc, "6.2", "打开共享列表", "点击“共享给我”，查看仍在有效期内的文件、所有者邮箱和到期时间。过期授权不会出现在有效列表中。")
    add_step(doc, "6.3", "下载并保存", "点击“下载”。系统通过权限验证后完成解密和完整性复核，浏览器按照原始文件名保存文件。下载次数会累加，便于所有者评估共享使用情况。")

    doc.add_heading("7  查看安全审计日志", level=1)
    doc.add_paragraph("点击“安全日志”可以查看本人操作以及自己所拥有文件的相关活动。日志包含操作类型、操作账号、结果、说明和时间，适合在答辩中证明系统具有可追溯性。")
    add_table(doc, ["日志类型", "触发条件", "建议关注点"], [
        ("加密上传", "文件成功加密并保存", "文件名与时间是否正确"),
        ("授权共享", "所有者创建或更新权限", "收件邮箱与有效期"),
        ("撤销授权", "所有者提前取消共享", "收件人与撤销时间"),
        ("解密下载", "权限和完整性检查通过", "下载者与操作时间"),
        ("删除文件", "所有者确认删除", "是否为预期文件"),
        ("操作失败", "上传、共享、下载、撤销或删除异常", "联系管理员查看服务日志"),
    ], [1.45, 2.35, 2.45])

    doc.add_heading("8  加密机制说明", level=1)
    doc.add_paragraph("系统采用混合加密，不使用 RSA 直接加密大文件。AES 适合高速处理任意文件内容，RSA 适合保护体积很小的 AES 密钥。两者结合可以兼顾性能、完整性和密钥管理。")
    add_table(doc, ["阶段", "输入", "处理", "输出"], [
        ("生成密钥", "安全随机源", "生成 AES 256 位密钥和 96 位 IV", "文件独立密钥材料"),
        ("内容加密", "原始文件", "AES 256 GCM 加密并产生认证标签", "仅可存储的密文"),
        ("密钥封装", "AES 密钥", "RSA OAEP SHA 256 公钥封装", "封装后的文件密钥"),
        ("完整性摘要", "原始文件", "计算 SHA 256", "64 位十六进制摘要"),
        ("下载解密", "密文与封装密钥", "鉴权 解封 解密 摘要复核", "验证后的原始文件"),
    ], [1.05, 1.45, 2.45, 1.3])
    doc.add_paragraph("RSA 私钥只存放在服务端受保护的环境变量中，不进入浏览器、数据库或源代码仓库。网站使用 HTTPS 提供传输层保护；应用层加密用于保护存储中的文件内容。")

    doc.add_heading("9  常见问题", level=1)
    faq = [
        ("看不到别人共享的文件", "确认登录邮箱与对方填写的邮箱完全一致，并确认授权未过期。"),
        ("上传按钮一直显示处理中", "检查网络连接，等待当前请求结束。刷新前不要重复提交同一文件。"),
        ("提示完整性校验失败", "不要继续使用该文件。该提示说明密文、认证标签或元数据不匹配，应联系管理员检查存储和日志。"),
        ("下载后文件名乱码", "使用最新版浏览器。系统通过 UTF 8 文件名参数返回中文名称。"),
        ("能否上传超大文件", "当前演示版限制单文件 15 MB、单账户 5 GB。论文扩展可采用分块加密、断点续传和分块认证标签。"),
        ("如何提前停止共享", "进入“我的文件”，在“当前有效授权”中找到对应收件人，点击“撤销”并确认。"),
        ("删除后能否恢复", "不能。删除会同时清理密文和共享权限，重要文件应保留本地备份。"),
    ]
    for question, answer in faq:
        heading = doc.add_heading(question, level=3)
        set_keep_with_next(heading)
        doc.add_paragraph(answer)

    doc.add_heading("10  答辩演示建议", level=1)
    doc.add_paragraph("建议准备两个测试账号和一个小型 PDF 或文本文件。完整演示控制在 5 至 7 分钟，并提前清理无关文件。")
    add_table(doc, ["顺序", "演示动作", "要说明的技术点"], [
        ("1", "账号 A 登录并上传文件", "随机 AES 密钥 GCM 加密 RSA 封装"),
        ("2", "展示文件摘要和安全日志", "完整性和可追溯性"),
        ("3", "账号 A 共享给账号 B", "邮箱绑定和有限期授权"),
        ("4", "账号 B 登录并下载", "服务端鉴权 解封 解密 摘要复核"),
        ("5", "账号 A 查看下载次数并撤销", "实时权限控制和撤销审计"),
        ("6", "账号 B 再次尝试下载", "已撤销权限立即被拒绝"),
        ("7", "删除文件并刷新列表", "密文 元数据 共享权限联动清理"),
    ], [0.65, 2.45, 3.15])
    doc.add_paragraph("答辩时应明确说明：系统没有自行设计密码算法，而是正确组合成熟的标准算法。创新点主要体现在完整安全流程、文件级独立密钥、细粒度共享权限、双重完整性验证和审计闭环。")

    doc.add_heading("11  管理员部署说明", level=1)
    doc.add_paragraph("管理员部署前需要准备 D1 数据库、R2 对象存储以及一组 2048 位 RSA OAEP JWK 密钥。生产环境变量 RSA PUBLIC JWK 和 RSA PRIVATE JWK 必须成对配置，其中私钥应标记为敏感值。")
    add_table(doc, ["配置项", "用途", "安全要求"], [
        ("DB", "文件元数据 共享和审计日志", "仅服务端绑定"),
        ("FILES", "保存 AES GCM 密文", "禁止公开读取"),
        ("RSA PUBLIC JWK", "封装每个文件的 AES 密钥", "可读但不得随意更换"),
        ("RSA PRIVATE JWK", "下载时解封 AES 密钥", "必须作为秘密保存和备份"),
        ("SITE URL", "生成网站社交预览绝对地址", "使用正式 HTTPS 地址"),
    ], [1.55, 2.35, 2.35], compact=True, trailing_space=False)
    doc.add_paragraph("更换 RSA 私钥会导致旧文件无法解密，因此必须规划密钥版本，并保留仍被文件引用的历史私钥。")
    doc.add_paragraph("部署后访问健康检查接口，并用测试账号完成上传、共享、撤销和下载闭环，确认各项配置生效。")

    doc.core_properties.title = "密盾传输系统使用手册"
    doc.core_properties.subject = "基于 AES 256 GCM 与 RSA OAEP 的文件安全传输系统用户操作手册"
    doc.core_properties.author = "密盾传输项目组"
    doc.core_properties.keywords = "文件安全传输 AES GCM RSA OAEP 使用手册"
    doc.save(OUTPUT)
    ROOT_COPY.write_bytes(OUTPUT.read_bytes())
    print(str(OUTPUT))
    print(str(ROOT_COPY))


if __name__ == "__main__":
    build_document()
